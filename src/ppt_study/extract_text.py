from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER


def extract_pptx_text(path: Path) -> list[dict]:
    if path.suffix.lower() != ".pptx":
        return []
    try:
        prs = Presentation(str(path))
    except Exception:
        return []
    pages: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        bodies: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text)
            if not text:
                continue
            is_title = False
            if shape.is_placeholder:
                ph = shape.placeholder_format.type
                is_title = ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            if is_title and not title:
                title = text
            else:
                bodies.append(text)
        notes = ""
        if slide.has_notes_slide:
            notes = "\n".join(
                p.text for p in slide.notes_slide.notes_text_frame.paragraphs if p.text
            )
        pages.append(
            {
                "index": i,
                "title": title.strip(),
                "body": "\n".join(bodies).strip(),
                "notes": notes.strip(),
            }
        )
    return pages


def _split_pdf_page_text(raw: str) -> tuple[str, str]:
    lines = [
        line.strip()
        for line in str(raw or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
        if line.strip()
    ]
    if not lines:
        return "", ""
    return lines[0], "\n".join(lines[1:])


def extract_pdf_text(path: Path) -> list[dict]:
    if path.suffix.lower() != ".pdf":
        return []
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return []
    try:
        pdf = pdfium.PdfDocument(str(path))
    except Exception:
        return []
    pages: list[dict] = []
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            raw = ""
            try:
                textpage = page.get_textpage()
                try:
                    raw = textpage.get_text_bounded() or ""
                finally:
                    textpage.close()
            except Exception:
                raw = ""
            finally:
                page.close()
            title, body = _split_pdf_page_text(raw)
            pages.append(
                {
                    "index": i + 1,
                    "title": title,
                    "body": body,
                    "notes": "",
                }
            )
    finally:
        pdf.close()
    return pages
