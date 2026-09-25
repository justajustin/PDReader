from pathlib import Path

from pptx import Presentation

from ppt_study.extract_text import extract_pdf_text, extract_pptx_text


def test_extracts_title_body_notes(tmp_path: Path):
    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Sharpening revisited"
    slide.placeholders[1].text = "original - smoothed = detail"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "speaker notes"
    file_path = tmp_path / "demo.pptx"
    prs.save(file_path)

    pages = extract_pptx_text(file_path)
    assert len(pages) == 1
    assert pages[0]["index"] == 1
    assert "Sharpening" in pages[0]["title"]
    assert "detail" in pages[0]["body"]
    assert "speaker notes" in pages[0]["notes"]


def test_ppt_binary_returns_empty(tmp_path: Path):
    path = tmp_path / "old.ppt"
    path.write_bytes(b"not a zip")
    assert extract_pptx_text(path) == []


def test_extract_pdf_text_reads_hello_page(tmp_path: Path):
    from pdf_util import write_hello_pdf

    path = write_hello_pdf(tmp_path / "hello.pdf", "Hello PDF")
    pages = extract_pdf_text(path)
    assert len(pages) == 1
    assert pages[0]["index"] == 1
    assert "Hello" in pages[0]["title"]
    assert pages[0]["notes"] == ""


def test_extract_pdf_text_rejects_non_pdf(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("x", encoding="utf-8")
    assert extract_pdf_text(path) == []


def test_extract_pdf_text_bad_file_returns_empty(tmp_path: Path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    assert extract_pdf_text(path) == []
