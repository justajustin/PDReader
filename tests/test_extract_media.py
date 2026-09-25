from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

from ppt_study.extract_media import extract_slide_media, publish_slide_media


def _pptx_with_video(folder: Path) -> Path:
    poster = folder / "poster.png"
    Image.new("RGB", (16, 12), "green").save(poster)
    clip = folder / "clip.webm"
    clip.write_bytes(b"FAKEWEBM")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_movie(
        str(clip),
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(3),
        poster_frame_image=str(poster),
        mime_type="video/webm",
    )
    dest = folder / "with-video.pptx"
    prs.save(dest)
    return dest


def test_extract_slide_media_finds_webm(tmp_path: Path):
    ppt = _pptx_with_video(tmp_path)
    items = extract_slide_media(ppt, 1)
    assert len(items) == 1
    item = items[0]
    assert item["kind"] == "video"
    assert item["ext"] == ".webm"
    assert item["blob"] == b"FAKEWEBM"
    assert 0 <= item["left"] < 0.5
    assert 0 < item["width"] < 1
    assert 0 < item["height"] < 1


def test_extract_slide_media_finds_gif(tmp_path: Path):
    gif = tmp_path / "spin.gif"
    gif.write_bytes(
        bytes.fromhex(
            "47494638396101000100800000ffffff00000021f90401000000002c000000000100010000020144003b"
        )
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(gif), Emu(0), Emu(0), Inches(2), Inches(2))
    dest = tmp_path / "with-gif.pptx"
    prs.save(dest)
    items = extract_slide_media(dest, 1)
    assert any(i["kind"] == "gif" for i in items)


def test_extract_slide_media_empty_for_text_only(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "No media"
    dest = tmp_path / "plain.pptx"
    prs.save(dest)
    assert extract_slide_media(dest, 1) == []


def test_publish_slide_media_writes_ui_files(appdata_tmp, tmp_path: Path):
    ppt = _pptx_with_video(tmp_path)
    items = publish_slide_media(ppt, 1, "deckkey")
    assert items[0]["src"].startswith("media/deckkey/")
    assert items[0]["src"].endswith(".webm")
    dest = appdata_tmp / "PPTStudyCompanion" / "ui" / Path(items[0]["src"])
    assert dest.is_file()
    assert dest.read_bytes() == b"FAKEWEBM"


def test_extract_slide_links_finds_http_url(tmp_path: Path):
    from pptx.util import Inches

    from ppt_study.extract_media import extract_slide_links

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(0.8))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Bilinear interpolation"
    run.hyperlink.address = "http://en.wikipedia.org/wiki/Bilinear_interpolation"
    dest = tmp_path / "with-link.pptx"
    prs.save(dest)
    links = extract_slide_links(dest, 1)
    assert links
    assert links[0]["url"] == "http://en.wikipedia.org/wiki/Bilinear_interpolation"
    assert 0 <= links[0]["left"] < 1
    assert 0 < links[0]["width"] <= 1


def test_extract_slide_links_empty_without_hyperlink(tmp_path: Path):
    from ppt_study.extract_media import extract_slide_links

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "No link"
    dest = tmp_path / "plain.pptx"
    prs.save(dest)
    assert extract_slide_links(dest, 1) == []
