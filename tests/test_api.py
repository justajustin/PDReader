import time
from pathlib import Path

from PIL import Image

from ppt_study.api import BridgeApi
from ppt_study.models import DeckRecord, SlideRecord


def _png(path: Path):
    Image.new("RGB", (32, 24), "white").save(path)


def test_select_and_script_cache(appdata_tmp, tmp_path, monkeypatch):
    img = tmp_path / "s.png"
    thumb = tmp_path / "t.png"
    _png(img)
    _png(thumb)
    deck = DeckRecord(
        str(tmp_path / "a.pptx"),
        1.0,
        [SlideRecord(1, "T", "B", "", str(thumb), str(img), None)],
    )
    api = BridgeApi()
    api.deck = deck
    api.current_index = 1
    monkeypatch.setattr(
        "ppt_study.api.AIClient.chat",
        lambda self, messages: "这是讲稿",
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    selected = api.select_slide(1)
    assert selected["ok"] is True
    assert selected["slide"]["title"] == "T"
    assert selected["slide"]["display_url"].startswith("data:image")
    assert selected["slide"]["media"] == []
    assert selected["slide"]["links"] == []
    out = api.generate_script()
    assert out["ok"] is True
    assert "讲稿" in out["script"]


def test_select_slide_returns_embedded_video(appdata_tmp, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    poster = tmp_path / "poster.png"
    Image.new("RGB", (16, 12), "green").save(poster)
    clip = tmp_path / "clip.webm"
    clip.write_bytes(b"FAKEWEBM")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_movie(
        str(clip),
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(3),
        poster_frame_image=str(poster),
        mime_type="video/webm",
    )
    ppt = tmp_path / "with-video.pptx"
    prs.save(ppt)
    img = tmp_path / "s.png"
    thumb = tmp_path / "t.png"
    _png(img)
    _png(thumb)
    api = BridgeApi()
    api.deck = DeckRecord(
        str(ppt),
        ppt.stat().st_mtime,
        [SlideRecord(1, "Wagon", "", "", str(thumb), str(img), None)],
    )
    out = api.select_slide(1)
    assert out["ok"] is True
    media = out["slide"]["media"]
    assert media
    assert media[0]["kind"] == "video"
    assert media[0]["src"].endswith(".webm")
    assert 0 <= media[0]["left"] < 1
    assert 0 < media[0]["width"] <= 1


def test_select_slide_returns_hyperlinks(appdata_tmp, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.6))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "wiki"
    run.hyperlink.address = "https://en.wikipedia.org/wiki/Bilinear_interpolation"
    ppt = tmp_path / "link.pptx"
    prs.save(ppt)
    img = tmp_path / "s.png"
    thumb = tmp_path / "t.png"
    _png(img)
    _png(thumb)
    api = BridgeApi()
    api.deck = DeckRecord(
        str(ppt),
        ppt.stat().st_mtime,
        [SlideRecord(1, "Filters", "", "", str(thumb), str(img), None)],
    )
    out = api.select_slide(1)
    assert out["ok"] is True
    links = out["slide"]["links"]
    assert links
    assert "Bilinear_interpolation" in links[0]["url"]


def test_open_url_only_allows_http(monkeypatch):
    seen = []
    monkeypatch.setattr("webbrowser.open", lambda url: seen.append(url) or True)
    api = BridgeApi()
    bad = api.open_url("javascript:alert(1)")
    assert bad["ok"] is False
    assert seen == []
    ok = api.open_url("https://en.wikipedia.org/wiki/Bilinear_interpolation")
    assert ok["ok"] is True
    assert seen == ["https://en.wikipedia.org/wiki/Bilinear_interpolation"]


def test_ask_blocked_without_key(appdata_tmp, monkeypatch):
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "", "gpt-5.6-sol", "bearer"
        ),
    )
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "T", "B", "", "", "", None)],
    )
    api.current_index = 1
    out = api.ask_question("什么是锐化？")
    assert out["ok"] is False
    assert "API Key" in out["error"]


def test_choose_and_open_cancel(monkeypatch):
    from ppt_study.api import BridgeApi
    api = BridgeApi()
    api._window = type("W", (), {"create_file_dialog": staticmethod(lambda *a, **k: None)})()
    out = api.choose_and_open()
    assert out["ok"] is False


def test_choose_and_open_returns_before_ingest_finishes(monkeypatch):
    import threading
    import time

    entered = threading.Event()
    release = threading.Event()

    def slow_open(self, path):
        entered.set()
        release.wait(timeout=5)
        return {"ok": True, "error": "", "deck": {"file_name": "a.pptx", "slides": []}}

    monkeypatch.setattr("ppt_study.api.BridgeApi.open_deck", slow_open)
    api = BridgeApi()
    api._window = type(
        "W",
        (),
        {
            "create_file_dialog": staticmethod(lambda *a, **k: [r"C:\a.pptx"]),
            "evaluate_js": staticmethod(lambda *a, **k: None),
        },
    )()
    t0 = time.monotonic()
    out = api.choose_and_open()
    elapsed = time.monotonic() - t0
    assert out["ok"] is True
    assert out.get("pending") is True
    assert elapsed < 0.75
    assert entered.wait(timeout=2)
    release.set()


def test_ppt_path_from_drop_event_picks_pptx():
    from ppt_study.api import ppt_path_from_drop_event

    path = ppt_path_from_drop_event(
        {
            "dataTransfer": {
                "files": [
                    {"name": "note.txt", "pywebviewFullPath": r"C:\note.txt"},
                    {"name": "lec.pptx", "pywebviewFullPath": r"C:\lec.pptx"},
                ]
            }
        }
    )
    assert path.endswith("lec.pptx")


def test_ppt_path_from_drop_event_picks_pdf():
    from ppt_study.api import ppt_path_from_drop_event

    path = ppt_path_from_drop_event(
        {
            "dataTransfer": {
                "files": [
                    {"name": "note.txt", "pywebviewFullPath": r"C:\note.txt"},
                    {"name": "lec.pdf", "pywebviewFullPath": r"C:\lec.pdf"},
                ]
            }
        }
    )
    assert path.endswith("lec.pdf")


def test_ppt_path_from_drop_event_rejects_other_files():
    from ppt_study.api import ppt_path_from_drop_event

    assert (
        ppt_path_from_drop_event(
            {
                "dataTransfer": {
                    "files": [{"name": "a.png", "pywebviewFullPath": r"C:\a.png"}]
                }
            }
        )
        == ""
    )


def test_open_from_path_rejects_non_ppt(tmp_path):
    api = BridgeApi()
    other = tmp_path / "notes.txt"
    other.write_text("x", encoding="utf-8")
    out = api.open_from_path(str(other))
    assert out["ok"] is False
    assert "PDF" in out["error"]


def test_open_from_path_accepts_pdf_suffix(tmp_path, monkeypatch):
    import threading

    pdf = tmp_path / "demo.pdf"
    pdf.write_bytes(b"%PDF")
    entered = threading.Event()
    release = threading.Event()

    def slow_open(self, path):
        entered.set()
        release.wait(timeout=5)
        return {"ok": True, "error": "", "deck": {"file_name": "demo.pdf", "slides": []}}

    monkeypatch.setattr("ppt_study.api.BridgeApi.open_deck", slow_open)
    api = BridgeApi()
    api._window = type("W", (), {"evaluate_js": staticmethod(lambda *a, **k: None)})()
    out = api.open_from_path(str(pdf))
    assert out["ok"] is True
    assert out.get("pending") is True
    assert entered.wait(timeout=2)
    release.set()


def test_choose_and_open_file_types_include_pdf():
    captured = {}

    def create_file_dialog(*args, **kwargs):
        captured["file_types"] = kwargs.get("file_types") or (args[1] if len(args) > 1 else ())
        return None

    api = BridgeApi()
    api._window = type("W", (), {"create_file_dialog": staticmethod(create_file_dialog)})()
    api.choose_and_open()
    blob = " ".join(captured.get("file_types") or ()).lower()
    assert "pdf" in blob
    assert "ppt" in blob


def test_open_from_path_rejects_missing(tmp_path):
    api = BridgeApi()
    out = api.open_from_path(str(tmp_path / "missing.pptx"))
    assert out["ok"] is False
    assert "找不到" in out["error"]


def test_open_from_path_returns_before_ingest_finishes(tmp_path, monkeypatch):
    import threading
    import time

    ppt = tmp_path / "demo.pptx"
    ppt.write_bytes(b"PK")
    entered = threading.Event()
    release = threading.Event()

    def slow_open(self, path):
        entered.set()
        release.wait(timeout=5)
        return {"ok": True, "error": "", "deck": {"file_name": "demo.pptx", "slides": []}}

    monkeypatch.setattr("ppt_study.api.BridgeApi.open_deck", slow_open)
    api = BridgeApi()
    api._window = type("W", (), {"evaluate_js": staticmethod(lambda *a, **k: None)})()
    t0 = time.monotonic()
    out = api.open_from_path(str(ppt))
    elapsed = time.monotonic() - t0
    assert out["ok"] is True
    assert out.get("pending") is True
    assert elapsed < 0.75
    assert entered.wait(timeout=2)
    release.set()


def test_js_api_does_not_expose_webview_window():
    api = BridgeApi()
    public_names = [name for name in dir(api) if not name.startswith("_")]
    assert "window" not in public_names


def test_save_settings_accepts_json_string(appdata_tmp):
    api = BridgeApi()
    out = api.save_settings(
        '{"base_url":"https://example.test/v1","api_key":"sk-new","model":"m","auth_mode":"raw"}'
    )
    assert out["ok"] is True
    from ppt_study.settings import load_settings

    loaded = load_settings()
    assert loaded.api_key == "sk-new"
    assert loaded.auth_mode == "raw"


def test_save_settings_persists_theme(appdata_tmp):
    api = BridgeApi()
    out = api.save_settings(
        '{"base_url":"https://example.test/v1","api_key":"sk-new","model":"m","auth_mode":"raw","theme":"dark"}'
    )
    assert out["ok"] is True
    from ppt_study.settings import load_settings

    assert load_settings().theme == "dark"
    got = api.get_settings()
    assert got["theme"] == "dark"


def test_save_settings_keeps_empty_api_fields(appdata_tmp):
    api = BridgeApi()
    out = api.save_settings({"base_url": "", "model": "", "theme": "light"})
    assert out["ok"] is True
    loaded = api.get_settings()
    assert loaded["base_url"] == ""
    assert loaded["model"] == ""
    again = api.save_settings({"theme": "dark"})
    assert again["ok"] is True
    kept = api.get_settings()
    assert kept["base_url"] == ""
    assert kept["model"] == ""
    assert kept["theme"] == "dark"


def test_generate_script_wraps_unexpected_errors(appdata_tmp, monkeypatch):
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "T", "B", "", "", "", None)],
    )
    api.current_index = 1
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr(
        "ppt_study.api.AIClient.chat",
        lambda self, messages: (_ for _ in ()).throw(IndexError("list index out of range")),
    )
    out = api.generate_script()
    assert out["ok"] is False
    assert "list index out of range" not in (out.get("error") or "")
    assert "生成失败" in out["error"] or "响应" in out["error"]


def test_generate_script_retries_timeout_once(appdata_tmp, monkeypatch):
    from ppt_study.ai_client import AIError
    from ppt_study.api import retryable_script_error

    assert retryable_script_error("积分服务器超时") is True
    assert retryable_script_error("请求超时") is True
    assert retryable_script_error("接口错误 503：Our servers are currently overloaded.") is True
    assert retryable_script_error("接口错误503：overloaded") is True
    assert retryable_script_error("积分不足，请充值") is False

    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "T", "B", "", "", "", None)],
    )
    api.current_index = 1
    monkeypatch.setattr("ppt_study.api.SCRIPT_RETRY_DELAY", 0)
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    n = {"i": 0}

    def fake_chat(self, messages):
        n["i"] += 1
        if n["i"] == 1:
            raise AIError("请求超时")
        return "重试后的讲稿"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    out = api.generate_script()
    assert out["ok"] is True
    assert out["script"] == "重试后的讲稿"
    assert n["i"] == 2


def test_generate_script_retries_overload_503(appdata_tmp, monkeypatch):
    from ppt_study.ai_client import AIError

    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "T", "B", "", "", "", None)],
    )
    api.current_index = 1
    monkeypatch.setattr("ppt_study.api.SCRIPT_RETRY_DELAY", 0)
    monkeypatch.setattr("ppt_study.api.SCRIPT_OVERLOAD_DELAY", 0)
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    n = {"i": 0}

    def fake_chat(self, messages):
        n["i"] += 1
        if n["i"] < 3:
            raise AIError("接口错误 503：Our servers are currently overloaded. Please try again later.")
        return "过载后恢复"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    out = api.generate_script()
    assert out["ok"] is True
    assert out["script"] == "过载后恢复"
    assert n["i"] == 3


def test_generate_script_does_not_retry_insufficient_credits(appdata_tmp, monkeypatch):
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "T", "B", "", "", "", None)],
    )
    api.current_index = 1
    monkeypatch.setattr("ppt_study.api.SCRIPT_RETRY_DELAY", 0)
    calls = {"n": 0}

    def fake_remote(*_a, **_k):
        calls["n"] += 1
        return {"ok": False, "error": "积分不足，请充值或改用自有 API"}

    monkeypatch.setattr("ppt_study.api.complete_remote", fake_remote)
    out = api.generate_script("platform")
    assert out["ok"] is False
    assert "积分不足" in (out.get("error") or "")
    assert calls["n"] == 1


def test_generate_script_streams_chunks(appdata_tmp, tmp_path, monkeypatch):
    from ppt_study.ai_client import StreamChunk

    img = tmp_path / "s.png"
    _png(img)
    api = BridgeApi()
    api.deck = DeckRecord(
        str(tmp_path / "a.pptx"),
        1.0,
        [SlideRecord(1, "T", "B", "", "", str(img), None)],
    )
    api.current_index = 1

    class FakeWindow:
        def __init__(self):
            self.calls = []

        def evaluate_js(self, code):
            self.calls.append(code)

    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)

    def fake_iter(self, messages, web_search=False):
        yield StreamChunk(text="本")
        yield StreamChunk(text="页讲稿")
        yield StreamChunk(done=True)

    monkeypatch.setattr("ppt_study.api.AIClient.iter_complete", fake_iter)
    api._window = FakeWindow()
    out = api.generate_script()
    assert out["ok"] is True
    assert out["pending"] is True
    blob = ""
    for _ in range(80):
        blob = "\n".join(api._window.calls)
        if "onScriptDone" in blob:
            break
        time.sleep(0.05)
    assert "onScriptChunk" in blob
    assert "本页讲稿" in blob
    assert "onScriptDone" in blob


def test_ask_question_streams_chunks(appdata_tmp, tmp_path, monkeypatch):
    from ppt_study.ai_client import StreamChunk

    img = tmp_path / "s.png"
    _png(img)
    api = BridgeApi()
    api.deck = DeckRecord(
        str(tmp_path / "a.pptx"),
        1.0,
        [SlideRecord(1, "二次型", "矩阵", "", "", str(img), None)],
    )
    api.current_index = 1

    class FakeWindow:
        def __init__(self):
            self.calls = []

        def evaluate_js(self, code):
            self.calls.append(code)

    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )

    def fake_iter(self, messages, web_search=False):
        yield StreamChunk(text="二次型")
        yield StreamChunk(text="来自矩阵")
        yield StreamChunk(done=True)

    monkeypatch.setattr("ppt_study.api.AIClient.iter_complete", fake_iter)
    api._window = FakeWindow()
    out = api.ask_question("这个矩阵二次型是怎么得到的")
    assert out["ok"] is True
    assert out["pending"] is True
    blob = ""
    for _ in range(80):
        blob = "\n".join(api._window.calls)
        if "onQaDone" in blob:
            break
        time.sleep(0.05)
    assert "onQaChunk" in blob
    assert "二次型来自矩阵" in blob
    assert "onQaDone" in blob


def test_generate_all_scripts_fills_missing_and_skips_existing(appdata_tmp, monkeypatch):
    monkeypatch.setattr("ppt_study.api.SCRIPT_BATCH_WORKERS", 1)
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [
            SlideRecord(1, "A", "B", "", "", "", "已有讲稿"),
            SlideRecord(2, "C", "D", "", "", "", None),
            SlideRecord(3, "E", "F", "", "", "", None),
        ],
    )
    api.current_index = 1
    calls = []
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)

    def fake_chat(self, messages):
        calls.append(1)
        return f"新讲稿{len(calls)}"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    out = api.generate_all_scripts()
    assert out["ok"] is True
    assert out["total"] == 3
    assert out["skipped"] == 1
    assert out["generated"] == 2
    assert api.deck.slides[0].script == "已有讲稿"
    assert api.deck.slides[1].script == "新讲稿1"
    assert api.deck.slides[2].script == "新讲稿2"
    assert len(calls) == 2


def test_generate_all_scripts_stops_on_error(appdata_tmp, monkeypatch):
    monkeypatch.setattr("ppt_study.api.SCRIPT_BATCH_WORKERS", 1)
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [
            SlideRecord(1, "A", "", "", "", "", None),
            SlideRecord(2, "B", "", "", "", "", None),
        ],
    )
    api.current_index = 1
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    n = {"i": 0}

    def fake_chat(self, messages):
        n["i"] += 1
        if n["i"] > 1:
            raise RuntimeError("积分不足，请充值或改用自有 API")
        return "第一页"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    out = api.generate_all_scripts()
    assert out["ok"] is False
    assert out["generated"] == 1
    assert "积分不足" in (out.get("error") or "") or "生成失败" in (out.get("error") or "")
    assert api.deck.slides[0].script == "第一页"
    assert not (api.deck.slides[1].script or "").strip()


def test_generate_all_scripts_continues_after_overload(appdata_tmp, monkeypatch):
    from ppt_study.ai_client import AIError

    monkeypatch.setattr("ppt_study.api.SCRIPT_BATCH_WORKERS", 1)
    monkeypatch.setattr("ppt_study.api.SCRIPT_RETRY_DELAY", 0)
    monkeypatch.setattr("ppt_study.api.SCRIPT_OVERLOAD_DELAY", 0)
    monkeypatch.setattr("ppt_study.api.SCRIPT_OVERLOAD_RETRIES", 1)
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [
            SlideRecord(1, "A", "", "", "", "", None),
            SlideRecord(2, "B", "", "", "", "", None),
            SlideRecord(3, "C", "", "", "", "", None),
        ],
    )
    api.current_index = 1
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    n = {"i": 0}

    def fake_chat(self, messages):
        n["i"] += 1
        if n["i"] in {2, 3}:
            raise AIError("接口错误503：Our servers are currently overloaded.")
        return "讲稿" + str(n["i"])

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    out = api.generate_all_scripts()
    assert out["ok"] is False
    assert out["generated"] == 2
    assert out["failed"] == 1
    assert "503" in (out.get("error") or "")
    assert api.deck.slides[0].script == "讲稿1"
    assert not (api.deck.slides[1].script or "").strip()
    assert api.deck.slides[2].script == "讲稿4"


def test_generate_all_scripts_emits_retry_progress(appdata_tmp, monkeypatch):
    from ppt_study.ai_client import AIError

    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "A", "", "", "", "", None)],
    )
    api.current_index = 1
    monkeypatch.setattr("ppt_study.api.SCRIPT_RETRY_DELAY", 0)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    n = {"i": 0}

    def fake_chat(self, messages):
        n["i"] += 1
        if n["i"] == 1:
            raise AIError("积分服务器超时")
        return "重试后的讲稿"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    events = []
    out = api._generate_all_scripts_now("", [1], events.append)
    assert out["ok"] is True
    assert out["generated"] == 1
    assert n["i"] == 2
    assert any(item.get("status") == "retrying" for item in events)


def test_generate_all_scripts_requires_deck():
    api = BridgeApi()
    out = api.generate_all_scripts()
    assert out["ok"] is False
    assert "尚未打开课件" in (out.get("error") or "")


def test_generate_all_scripts_starts_next_before_first_finishes(appdata_tmp, monkeypatch):
    import threading

    monkeypatch.setattr("ppt_study.api.SCRIPT_BATCH_WORKERS", 2)
    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [
            SlideRecord(1, "A", "", "", "", "", None),
            SlideRecord(2, "B", "", "", "", "", None),
            SlideRecord(3, "C", "", "", "", "", None),
        ],
    )
    api.current_index = 1
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    barrier = threading.Barrier(2, timeout=2)
    inflight = {"n": 0, "max": 0, "calls": 0}
    lock = threading.Lock()

    def fake_chat(self, messages):
        with lock:
            inflight["calls"] += 1
            call = inflight["calls"]
            inflight["n"] += 1
            inflight["max"] = max(inflight["max"], inflight["n"])
        if call <= 2:
            barrier.wait()
        time.sleep(0.02)
        with lock:
            inflight["n"] -= 1
        return "讲稿"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    events = []
    out = api._generate_all_scripts_now("", [1, 2, 3], events.append)
    assert out["ok"] is True
    assert out["generated"] == 3
    assert inflight["max"] == 2
    running = [item for item in events if item.get("status") == "running"]
    assert any(len(item.get("active") or []) > 1 for item in running)


def test_generate_all_scripts_runs_in_background_thread(appdata_tmp, monkeypatch):
    import threading

    from ppt_study.ai_client import StreamChunk

    api = BridgeApi()
    api.deck = DeckRecord(
        "a.pptx",
        1.0,
        [SlideRecord(1, "A", "", "", "", "", None)],
    )
    api.current_index = 1
    started = threading.Event()
    caller = {"id": None}

    class DummyWindow:
        def evaluate_js(self, _code):
            pass

    def fake_iter(self, messages, web_search=False):
        caller["id"] = threading.current_thread().ident
        started.set()
        yield StreamChunk(text="后台讲稿")
        yield StreamChunk(done=True)

    monkeypatch.setattr("ppt_study.api.AIClient.iter_complete", fake_iter)
    monkeypatch.setattr("ppt_study.api.save_script", lambda *a, **k: None)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    api._window = DummyWindow()
    main_id = threading.current_thread().ident
    out = api.generate_all_scripts()
    assert out["pending"] is True
    assert started.wait(2)
    deadline = time.time() + 2
    while api.deck.slides[0].script != "后台讲稿" and time.time() < deadline:
        time.sleep(0.01)
    assert api.deck.slides[0].script == "后台讲稿"
    assert caller["id"] not in (None, main_id)


def test_ask_question_forwards_attachments(appdata_tmp, tmp_path, monkeypatch):
    img = tmp_path / "s.png"
    _png(img)
    captured = {}

    def fake_chat(self, messages):
        captured["messages"] = messages
        return "截图里是偏导数"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    monkeypatch.setattr(
        "ppt_study.api.data_url_to_api_data_url",
        lambda url: "data:image/jpeg;base64,ATT",
    )
    api = BridgeApi()
    api.deck = DeckRecord(
        str(tmp_path / "a.pptx"),
        1.0,
        [SlideRecord(1, "T", "B", "", "", str(img), None)],
    )
    api.current_index = 1
    out = api.ask_question("这个公式", ["data:image/png;base64,xxx"])
    assert out["ok"] is True
    last = captured["messages"][-1]
    parts = last["content"]
    urls = [p["image_url"]["url"] for p in parts if p.get("type") == "image_url"]
    assert "data:image/jpeg;base64,ATT" in urls


def _ready_api(monkeypatch, tmp_path):
    img = tmp_path / "s.png"
    _png(img)
    monkeypatch.setattr(
        "ppt_study.api.load_settings",
        lambda: __import__("ppt_study.models", fromlist=["AppSettings"]).AppSettings(
            "https://www.dmxapi.cn/v1", "sk-x", "gpt-5.6-sol", "bearer"
        ),
    )
    api = BridgeApi()
    api.deck = DeckRecord(
        str(tmp_path / "a.pptx"),
        1.0,
        [SlideRecord(1, "T", "B", "", "", str(img), None)],
    )
    api.current_index = 1
    return api


def test_clear_chat_empties_history(appdata_tmp):
    api = BridgeApi()
    api.history = [
        {"role": "user", "content": "q1"},
        {"role": "assistant", "content": "a1"},
    ]
    out = api.clear_chat()
    assert out["ok"] is True
    assert api.history == []


def test_delete_turn_removes_qa_pair(appdata_tmp):
    api = BridgeApi()
    api.history = [
        {"role": "user", "content": "q1"},
        {"role": "assistant", "content": "a1"},
        {"role": "user", "content": "q2"},
        {"role": "assistant", "content": "a2"},
    ]
    out = api.delete_turn(0)
    assert out["ok"] is True
    assert [m["content"] for m in api.history] == ["q2", "a2"]
    missing = api.delete_turn(4)
    assert missing["ok"] is False
    assert "问答" in missing["error"]


def test_retry_turn_replaces_answer_using_prior_context(appdata_tmp, tmp_path, monkeypatch):
    captured = {}

    def fake_chat(self, messages):
        captured["messages"] = messages
        return "新回答"

    monkeypatch.setattr("ppt_study.api.AIClient.chat", fake_chat)
    api = _ready_api(monkeypatch, tmp_path)
    api.history = [
        {"role": "user", "content": "第一问", "question": "第一问"},
        {"role": "assistant", "content": "旧1"},
        {"role": "user", "content": "第二问", "question": "第二问"},
        {"role": "assistant", "content": "旧2"},
    ]
    out = api.retry_turn(1)
    assert out["ok"] is True
    assert out["answer"] == "新回答"
    assert api.history[1]["content"] == "旧1"
    assert api.history[3]["content"] == "新回答"
    blob = []
    for msg in captured["messages"]:
        content = msg["content"]
        if isinstance(content, str):
            blob.append(content)
        elif isinstance(content, list):
            blob.extend(p.get("text", "") for p in content if isinstance(p, dict))
    text = "\n".join(blob)
    assert "第一问" in text
    assert "旧1" in text
    assert "第二问" in text
    assert "旧2" not in text


def test_open_deck_payload_skips_inline_thumbs(appdata_tmp, tmp_path, monkeypatch):
    thumb = tmp_path / "t.png"
    _png(thumb)
    deck = DeckRecord(
        str(tmp_path / "a.pptx"),
        1.0,
        [SlideRecord(1, "T", "B", "", str(thumb), str(thumb), None)],
    )
    monkeypatch.setattr("ppt_study.api.ingest_presentation", lambda *a, **k: deck)
    api = BridgeApi()
    out = api.open_deck(str(tmp_path / "a.pptx"))
    assert out["ok"] is True
    assert out["deck"]["slides"][0]["thumb_url"] == ""
    thumb_out = api.get_thumb(1)
    assert thumb_out["ok"] is True
    assert thumb_out["thumb_url"].startswith("data:image")


def test_open_deck_worker_streams_thumbs(appdata_tmp, tmp_path, monkeypatch):
    thumb = tmp_path / "t.png"
    _png(thumb)
    deck = DeckRecord(
        str(tmp_path / "b.pptx"),
        1.0,
        [SlideRecord(1, "T", "B", "", str(thumb), str(thumb), None)],
    )
    monkeypatch.setattr("ppt_study.api.ingest_presentation", lambda *a, **k: deck)
    calls: list[str] = []
    api = BridgeApi()
    api._window = type("W", (), {"evaluate_js": staticmethod(lambda code: calls.append(code))})()
    api._open_deck_worker(str(tmp_path / "b.pptx"))
    joined = "\n".join(calls)
    assert "onDeckReady" in joined
    assert "onThumbReady" in joined


def test_toggle_favorite_marks_open_deck_slides(appdata_tmp, tmp_path, monkeypatch):
    thumb = tmp_path / "t.png"
    _png(thumb)
    ppt = tmp_path / "lec.pptx"
    ppt.write_bytes(b"x")
    deck = DeckRecord(
        str(ppt),
        1.0,
        [
            SlideRecord(1, "A", "B", "", str(thumb), str(thumb), None),
            SlideRecord(2, "C", "D", "", str(thumb), str(thumb), None),
        ],
    )
    monkeypatch.setattr("ppt_study.api.ingest_presentation", lambda *a, **k: deck)
    api = BridgeApi()
    opened = api.open_deck(str(ppt))
    assert opened["deck"]["slides"][0]["starred"] is False
    out = api.toggle_favorite(1)
    assert out["ok"] is True
    assert out["starred"] is True
    opened = api.open_deck(str(ppt))
    stars = {s["index"]: s["starred"] for s in opened["deck"]["slides"]}
    assert stars[1] is True
    assert stars[2] is False


def test_qa_favorite_api_saves_and_searches(appdata_tmp):
    api = BridgeApi()
    saved = api.save_qa_favorite(
        {
            "question": "sinc 滤波器有什么特点",
            "answer": "理想但可能振铃。",
            "file_name": "lec03.pptx",
            "slide_index": 12,
        }
    )
    assert saved["ok"] is True
    assert saved["item"]["id"]
    listed = api.list_qa_favorites("振铃")
    assert listed["ok"] is True
    assert len(listed["items"]) == 1
    removed = api.remove_qa_favorite(saved["item"]["id"])
    assert removed["ok"] is True
    assert api.list_qa_favorites("")["items"] == []






def test_check_and_apply_app_update(appdata_tmp, tmp_path, monkeypatch):
    from ppt_study.paths import APP_VERSION

    launched = []

    def fake_info():
        return {
            "ok": True,
            "available": True,
            "version": "9.9.9",
            "size": 2,
            "sha256": "abc",
        }

    def fake_download(path, expected_sha256=""):
        assert expected_sha256 == "abc"
        dest = __import__("pathlib").Path(path)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(b"MZ")
        return {"ok": True, "path": str(dest), "error": ""}

    def fake_popen(args, **kwargs):
        launched.append(args)
        return type("P", (), {"pid": 1})()

    monkeypatch.setattr("ppt_study.api.fetch_update_info", fake_info)
    monkeypatch.setattr("ppt_study.api.download_update_file", fake_download)
    monkeypatch.setattr("ppt_study.api.subprocess.Popen", fake_popen)
    api = BridgeApi()
    checked = api.check_app_update()
    assert checked["ok"] is True
    assert checked["newer"] is True
    assert checked["current"] == APP_VERSION
    applied = api.apply_app_update()
    assert applied["ok"] is True
    assert applied["quit"] is True
    assert launched
    joined = " ".join(str(item) for item in launched[0])
    assert "/FORCECLOSEAPPLICATIONS" in joined
    assert "/APPVERSION=9.9.9" in joined
    assert "cmd.exe" in launched[0]
    monkeypatch.setattr(
        "ppt_study.api.fetch_update_info",
        lambda: {
            "ok": True,
            "available": True,
            "version": APP_VERSION,
            "size": 2,
            "sha256": "",
        },
    )
    same = api.check_app_update()
    assert same["newer"] is False
    blocked = api.apply_app_update()
    assert blocked["ok"] is False
    assert "最新" in (blocked.get("error") or "")


def test_check_app_update_uses_server_installed_version(appdata_tmp, monkeypatch):
    from ppt_study.app_version import save_installed_version

    save_installed_version("9.9.9")
    monkeypatch.setattr(
        "ppt_study.api.fetch_update_info",
        lambda: {
            "ok": True,
            "available": True,
            "version": "9.9.9",
            "size": 2,
            "sha256": "abc",
        },
    )
    api = BridgeApi()
    checked = api.check_app_update()
    assert checked["current"] == "9.9.9"
    assert checked["newer"] is False
    assert api.app_version()["version"] == "9.9.9"


def test_update_setup_command_delays_and_can_target_install_dir(monkeypatch, tmp_path):
    from ppt_study.api import update_setup_command

    setup = tmp_path / "PDReaderSetup.exe"
    cmd = update_setup_command(str(setup), "0.2.9")
    assert cmd[0] == "cmd.exe"
    assert "/FORCECLOSEAPPLICATIONS" in cmd[-1]
    assert "--silent" not in cmd[-1]
    monkeypatch.setattr("ppt_study.api.sys.frozen", True, raising=False)
    monkeypatch.setattr("ppt_study.api.sys.executable", str(tmp_path / "PDReader" / "PDReader.exe"))
    frozen = update_setup_command(str(setup), "0.2.9")
    assert "--silent" in frozen[-1]
    assert "--dir" in frozen[-1]


def test_quit_app_ready_for_update():
    api = BridgeApi()
    out = api.quit_app()
    assert out["ok"] is True
    assert out["quit"] is True


def test_hello_server_registers_without_redeem(monkeypatch):
    seen = {}

    def fake_hello():
        seen["called"] = True
        return {"ok": True, "account_id": "acc_x", "paid": False, "balance": "0", "error": ""}

    monkeypatch.setattr("ppt_study.api.hello_remote", fake_hello)
    api = BridgeApi()
    out = api.hello_server()
    assert seen["called"] is True
    assert out["ok"] is True
    assert out["paid"] is False
